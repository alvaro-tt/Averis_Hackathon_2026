import pdfplumber
import io
import openpyxl
from docx import Document

# Symbol fonts hold no real text: a PDF without a Chinese font renders
# "毛重" as ZapfDingbats boxes, which extract as a literal "nn".
_SYMBOL_FONTS = ("ZapfDingbats", "Symbol")


def _is_text_char(obj, bold):
    if obj.get("object_type") != "char":
        return True
    font = obj.get("fontname", "")
    if any(sym in font for sym in _SYMBOL_FONTS):
        return False
    return ("Bold" in font) == bold


def _font_aware_lines(page):
    """Rebuild lines keeping bold labels and regular values apart.

    Form-style PDFs draw a bold label at the left and a regular value in a
    column to its right. When a long label ("Notify Party/Intermediate
    Consignee") runs into the value column, plain extract_text() interleaves
    the two character by character. Extracting bold and regular text
    separately and re-joining by line fixes that, and emits "Label: value".
    """
    rows = []
    for bold in (True, False):
        sub = page.filter(lambda o, b=bold: _is_text_char(o, b))
        for line in sub.extract_text_lines():
            rows.append({"top": line["top"], "x0": line["x0"], "bold": bold, "text": line["text"]})
    rows.sort(key=lambda r: (r["top"], r["x0"]))

    grouped = []
    for row in rows:
        if grouped and abs(row["top"] - grouped[-1][0]["top"]) <= 2:
            grouped[-1].append(row)
        else:
            grouped.append([row])

    lines = []
    for group in grouped:
        group.sort(key=lambda r: r["x0"])
        if (len(group) == 2 and group[0]["bold"] and not group[1]["bold"]
                and not group[0]["text"].rstrip().endswith(":")):
            lines.append(f"{group[0]['text'].strip()}: {group[1]['text'].strip()}")
        else:
            lines.append(" ".join(r["text"].strip() for r in group))
    return lines


def read_pdf_text(file_bytes):
    """Extract all text from a PDF, page by page, joined with newlines.
    Uses font-aware line rebuilding when the page mixes bold and regular
    text; falls back to plain extract_text() otherwise or on any error.
    Returns "" for image-only (scanned) PDFs - there is no text layer."""
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            fonts = {c.get("fontname", "") for c in page.chars}
            mixed = any("Bold" in f for f in fonts) and any("Bold" not in f for f in fonts
                                                           if not any(s in f for s in _SYMBOL_FONTS))
            page_text = None
            if mixed:
                try:
                    page_text = "\n".join(_font_aware_lines(page))
                except Exception:
                    page_text = None
            if page_text is None:
                page_text = page.filter(
                    lambda o: o.get("object_type") != "char"
                    or not any(s in o.get("fontname", "") for s in _SYMBOL_FONTS)
                ).extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)

def _row_to_lines(cells):
    """One table/sheet row -> text lines.
    2 cells            -> "label: value"
    4 / 6 / 8 cells    -> label/value pairs side by side (common BL grid layout),
                          one "label: value" line per pair
    anything else      -> cells joined with spaces
    A cell that already contains "label: value" is kept as is."""
    if len(cells) == 2:
        return [f"{cells[0]}: {cells[1]}" if not cells[0].rstrip().endswith(":")
                else f"{cells[0]} {cells[1]}"]
    if len(cells) in (4, 6, 8) and not any(":" in c for c in cells):
        return [f"{cells[i]}: {cells[i + 1]}" for i in range(0, len(cells), 2)]
    return [" ".join(cells)]


def read_xlsx_text(file_bytes):
    """Extract all cell values from an XLSX, formatted as readable lines."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    text_parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                text_parts.extend(_row_to_lines(cells))
    return "\n".join(text_parts)


def _docx_table_lines(table):
    lines = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
        # merged cells repeat the same text: drop consecutive duplicates
        clean = []
        for cell in cells:
            if not clean or cell != clean[-1]:
                clean.append(cell)
        if clean:
            lines.extend(_row_to_lines(clean))
    return lines


def read_docx_text(file_bytes):
    """Extract paragraphs and tables from a DOCX in reading order (a label in
    a paragraph followed by its table stays together), joined with newlines."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(file_bytes))
    text_parts = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, doc).text.strip()
            if text:
                text_parts.append(text)
        elif tag == "tbl":
            text_parts.extend(_docx_table_lines(Table(child, doc)))
    return "\n".join(text_parts)